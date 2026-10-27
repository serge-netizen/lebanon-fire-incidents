import json

from openpyxl import load_workbook

from crawler import derive, review
from crawler.extract import PROMPT_VERSION
from crawler.store import Store

ITEM = {
    "record_type": "incident", "is_lebanon": True, "summary": "A fire.", "occurred_on": None,
    "date_precision": "unknown", "place_text": None, "site_primary": "dump", "site_others": [],
    "material_primary": "household_rubbish", "material_others": [], "hazard_primary": "not_established",
    "hazard_others": [], "origin": "unknown", "deaths": None, "injuries": None, "smoke_inhalation": None,
    "buildings_burned": None, "vehicles_burned": None, "area_burned_m2": None, "displaced": None,
    "responders": [], "evidence": [], "confidence": 0.9,
}
DUMP = ITEM | {"place_text": "صيدا - جبل النفايات", "evidence": [{"field": "site_primary", "quote": "جبل النفايات"}]}
ARTICLES = {  # key: (outlet, headline, items)
    "a": ("nna", "حريق جبل النفايات في صيدا", [DUMP]),
    "b": ("janoubia", "حرائق النفايات في صيدا", [DUMP | {"site_primary": "not_established", "injuries": 1}]),
    "c": ("nna", "غارة تشعل حريقاً في الخيام", [ITEM | {"place_text": "الخيام", "site_primary": "orchard",
                                                       "origin": "conflict_related"}]),
    "d": ("nna", "حريق في كنيسة في مصر", [ITEM | {"is_lebanon": False}]),
}


def test_workbook_has_one_row_per_fire(tmp_path):
    for key, (outlet, headline, items) in ARTICLES.items():
        (tmp_path / "raw" / outlet / "2026-10").mkdir(parents=True, exist_ok=True)
        (tmp_path / "extracted" / outlet / "2026-10").mkdir(parents=True, exist_ok=True)
        article = {"url": f"https://x.test/{key}", "outlet": outlet, "headline": headline,
                   "published_at": "2026-10-07T12:00:00+03:00", "text": headline}
        (tmp_path / f"raw/{outlet}/2026-10/{key}.json").write_text(json.dumps(article, ensure_ascii=False), encoding="utf-8")
        result = {"raw_ref": f"raw/{outlet}/2026-10/{key}.json", "prompt_version": PROMPT_VERSION, "status": "ok",
                  "items": items}
        (tmp_path / f"extracted/{outlet}/2026-10/{key}.json").write_text(json.dumps(result, ensure_ascii=False),
                                                                          encoding="utf-8")
    out = tmp_path / "review.xlsx"
    counts = review.write(out, Store(tmp_path))
    assert (counts["fires"], counts["war_related"], counts["merged"], counts["set_aside"]) == (1, 1, 1, 1)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Read me", "Fires in Lebanon", "War-related (not shown)", "Articles", "Set aside", "Lists"]
    lists = {row[1]: row for row in wb["Lists"].iter_rows(min_row=2, values_only=True)}
    assert lists["bakery"][5].startswith("what burned: Diesel and fuel oil") and "hazard: Cooking" in lists["restaurant"][5]
    assert lists["cooking_oil_fat"][5] == "hazard: Cooking" and lists["gas"][5] == "what burned: LPG and gas cylinders · origin: Accidental"
    assert lists["vehicle_other"][2].startswith("Other vehicle") and lists["4"][2] == "Deaths or several buildings"
    header = [c.value for c in wb["Fires in Lebanon"][1]]
    row = dict(zip(header, [c.value for c in wb["Fires in Lebanon"][2]], strict=True))
    assert row["Date"] == "2026-10-07\n— deduced from the publication date"
    assert row["What happened (Claude's summary; click to open the article)"].endswith("[2 articles: nna, janoubia]")
    assert row["Site that witnessed the fire"] == "Dump or landfill\n— reported"
    assert row["What burned"] == "Household rubbish\n— reported"
    assert row["Hazard"] == "Open burning; Natural causes\n— deduced from the site: usual at “Dump or landfill”"
    assert row["Damage"] == "1 injured" and row["Gravity of the damage"] == "3 — Injuries or displaced people"
    assert row["Origin"] == "not established" and row["Responders"] == "not stated"
    assert [c.value for c in wb["Articles"]["A"]][1:] == ["F1", "F1", "W1"]


def test_a_place_beside_the_fire_stays_reported_under_a_deduced_main_site():
    record = {"site_primary": "not_established", "site_others": ["education"], "material_primary": "dry_grass_weeds",
              "material_others": []}
    text = review._question(record, "site", derive.site_label, derive.site(record))
    assert text == ("Natural areas\n— deduced from what burned (Dry grass and weeds)\n"
                    "also: School, university or nursery — reported")


def test_a_cause_the_article_does_not_settle_is_shown_as_possible():
    record = {"hazard_primary": "not_established", "hazard_others": ["munitions_explosives", "vehicle_fault"]}
    text = review._question(record, "hazard", lambda c: c, None)
    assert text == "not established (possible: munitions_explosives; vehicle_fault)\n— reported"
