"""Review (BRIEF §4.7). For now, the owner's review workbook — the task 4 feedback loop: one row
per fire (articles about the same fire share a row), answering the owner's nine questions
(2026-10-27) — date, location, site, what burned, hazard, origin, damage, gravity, responders —
each marked reported or deduced, and two columns for the owner's verdict and correction.
Private: it quotes the articles, so it is written to the private store, never to this repo.
Task 5 adds the review sheet and reads decisions back.
"""

from __future__ import annotations

from collections import Counter
from collections.abc import Callable
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from crawler import derive
from crawler import taxonomy as tx
from crawler.dedupe import DAMAGES, Incident, incidents, load
from crawler.derive import WAR, Deduced
from crawler.geocode import Resolution
from crawler.store import Store

HEADER_FILL = PatternFill("solid", fgColor="2F3B4C")
REVIEW_FILL = PatternFill("solid", fgColor="FFF4D6")
WRAP = Alignment(wrap_text=True, vertical="top")
LINK = Font(color="1F5FAF", underline="single")
LOW_CONFIDENCE = 0.7  # below this, the brief sends a record to a second pass (BRIEF §4.4)
NOT_ESTABLISHED = "not established"

FIRE_COLUMNS = [  # (title, width): the owner's nine questions, after the summary they are judged by
    ("#", 5), ("Date", 13), ("Location", 26), ("What happened (Claude's summary; click to open the article)", 46),
    ("Site that witnessed the fire", 26), ("What burned", 28), ("Hazard", 26), ("Origin", 18), ("Damage", 22),
    ("Gravity of the damage", 18), ("Responders", 16), ("Checks", 30),
    ("YOUR VERDICT (OK / wrong)", 14), ("YOUR CORRECTION", 40),
]
ARTICLE_COLUMNS = [("Fire", 6), ("Outlet", 11), ("Published", 11), ("Headline", 50), ("Summary (Claude)", 50),
                   ("Evidence: Claude's quotes from the article", 70)]
SET_ASIDE_COLUMNS = [("#", 5), ("Published", 11), ("Outlet", 11), ("Headline", 60), ("Set aside as", 14),
                     ("Summary (Claude)", 60), ("Confidence", 10)]
LIST_COLUMNS = [("List", 14), ("Id", 22), ("Label", 34), ("Definition or rule", 50), ("Arabic words that point to it", 34),
                ("When reported, it lets us deduce", 50), ("YOUR COMMENT", 40)]
SET_ASIDE_AS = {"aggregate_report": "totals", "background": "background", "not_a_fire": "not a fire", "abroad": "abroad"}
DAMAGE_TEXT = {  # (one, several)
    "deaths": ("dead", "dead"), "injuries": ("injured", "injured"),
    "smoke_inhalation": ("treated for smoke", "treated for smoke"),
    "buildings_burned": ("building burned", "buildings burned"), "vehicles_burned": ("vehicle burned", "vehicles burned"),
    "area_burned_m2": ("m² burned", "m² burned"), "displaced": ("displaced", "displaced"),
}
DAMAGE_IN_WORDS = {"none": "no damage", "damaged": "damage stated, no figures"}


# ---------------------------------------------------------------- cell text

def _question(record: dict, kind: str, label: Callable[[str], str], deduced: Deduced | None) -> str:
    """One answer: reported (the main value, then the others), else deduced (with what it was
    deduced from), else not established."""
    primary, others = record[f"{kind}_primary"], [label(c) for c in record[f"{kind}_others"]]
    # other values; with no main hazard, they are causes the article names without settling
    unsettled = kind == "hazard" and primary == tx.NOT_ESTABLISHED
    named = f"{'possible' if unsettled else 'also'}: {'; '.join(others)}" if others else ""
    if primary != tx.NOT_ESTABLISHED:
        return label(primary) + (f" ({named})" if named else "") + "\n— reported"
    if deduced:
        return f"{'; '.join(label(c) for c in deduced.values)}\n— deduced {deduced.basis}" + (
            f"\n{named} — reported" if named else "")
    return f"{NOT_ESTABLISHED} ({named})\n— reported" if named else NOT_ESTABLISHED


def _location(r: Resolution, as_written: str | None) -> str:
    parts = [u["name_en"] for u in (r.admin3, r.admin2, r.admin1) if u]
    found = " · ".join(parts) if parts else "not found on the map"
    return f"{found}\nas written: {as_written}" if as_written else found


def _origin(record: dict) -> str:
    labels = tx.labels()["origin"]
    if record["origin"] != "unknown":
        return f"{labels[record['origin']]}\n— reported"
    deduced = derive.origin(record)
    return f"{labels[deduced.values[0]]}\n— deduced {deduced.basis}" if deduced else NOT_ESTABLISHED


def _damage(record: dict) -> str:
    parts = [f"{record[k]:,.0f} {DAMAGE_TEXT[k][record[k] != 1]}" for k in DAMAGES if record[k] is not None]
    return "; ".join(parts) or DAMAGE_IN_WORDS.get(record.get("property_damage", ""), "not stated")


def _gravity(record: dict) -> str:
    level = derive.gravity(record)
    return f"{level} — {tx.labels()['gravity'][level]}" if level else NOT_ESTABLISHED


def _outlets(i: Incident) -> str:
    return ", ".join(f"{o} ×{n}" if n > 1 else o for o, n in Counter(r.outlet for r in i.reports).items())


def fire_row(n: int, i: Incident) -> list:
    rec, main, labels = i.record, i.reports[0], tx.labels()
    count = f"{len(i.reports)} articles: " if len(i.reports) > 1 else ""
    checks = [*i.flags] + ([f"confidence below {LOW_CONFIDENCE}"] if rec["confidence"] < LOW_CONFIDENCE else [])
    date = i.date.isoformat() + ("\n— deduced from the publication date" if i.date_deduced else "\n— reported")
    return [
        n, date, _location(main.place, rec["place_text"]), (f"{rec['summary']}\n[{count}{_outlets(i)}]", main.url),
        _question(rec, "site", derive.site_label, derive.site(rec)),
        _question(rec, "material", labels["material"].__getitem__, derive.materials(rec)),
        _question(rec, "hazard", labels["hazard"].__getitem__, derive.hazards(rec)),
        _origin(rec), _damage(rec), _gravity(rec),
        ", ".join(labels["responder"][x] for x in rec["responders"]) or "not stated", "\n".join(checks), "", "",
    ]


def list_rows() -> list[list]:
    """Every entry of the lists, with what it lets us deduce, for the owner to correct or extend."""
    t, labels = tx.lists(), tx.labels()

    def names(kind: str, codes: list[str]) -> str:
        return "; ".join(derive.site_label(c) if kind == "site" else labels[kind][c] for c in codes)

    def deduce(*found: tuple[str, str, list[str]]) -> str:
        """(question, label kind, ids) → "what burned: …  ·  hazard: …"."""
        return " · ".join(f"{what}: {names(kind, codes)}" for what, kind, codes in found if codes)

    def text(entry: dict, *fields: str) -> str:
        return " ".join(str(entry[f]) for f in fields if entry.get(f))

    def typical(table: str, *keys: str) -> list[str]:
        return next((t[table][k] for k in keys if k in t[table]), [])

    rows = []
    for category, v in t["site_types"].items():
        rows.append(["Site", category, v["label"].upper(), text(v, "definition", "rule"), "",
                     deduce(("what burned", "material", typical("typical_materials", category)),
                            ("hazard", "hazard", typical("typical_hazards", category))), ""])
        rows += [["Site", sub, s["label"], text(s, "rule"), "، ".join(s.get("cues", [])),
                  deduce(("what burned", "material", typical("typical_materials", sub, category)),
                         ("hazard", "hazard", typical("typical_hazards", sub, category))), ""]
                 for sub, s in (v.get("subtypes") or {}).items()]
    for group in t["materials"].values():
        rows.append(["What burned", "", group["label"].upper(), text(group, "definition"), "", "", ""])
        rows += [["What burned", k, f"{m['label']} [class {m['fire_class']}]", text(m, "rule", "note"),
                  "، ".join(m.get("cues", [])),
                  deduce(("site", "site", [t["site_from_material"][k]] if k in t["site_from_material"] else []),
                         ("hazard", "hazard", [t["hazard_from_material"][k]] if k in t["hazard_from_material"] else [])),
                  ""]
                 for k, m in group["items"].items()]
    origin_of = {h: o for o, hazards in t["origin_from_hazard"].items() for h in hazards}
    rows += [["Hazard", k, h["label"], text(h, "definition", "rule"), "، ".join(h.get("cues", [])),
              deduce(("what burned", "material", [t["material_from_hazard"][k]] if k in t["material_from_hazard"] else []),
                     ("origin", "origin", [origin_of[k]] if k in origin_of else [])), ""]
             for k, h in t["hazards"].items()]
    rows += [["Origin", k, o["label"], o["definition"], "", "", ""] for k, o in t["origins"].items()]
    rows += [["Damage in words", k, d["label"], d["definition"], "", "", ""] for k, d in t["property_damage"].items()]
    rows += [["Gravity", str(k), g["label"], g["rule"], "", "", ""] for k, g in t["gravity"].items()]
    rows += [["Responders", k, r["label"], "", "، ".join(r["cues"]), "", ""] for k, r in t["responders"].items()]
    return rows


# ---------------------------------------------------------------- the workbook

def _sheet(wb: Workbook, title: str, header: list[tuple[str, int]], rows: list[list], review_cols: int = 0) -> None:
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.append([h for h, _ in header])
    for col, (_, width) in enumerate(header, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font, cell.fill, cell.alignment = Font(bold=True, color="FFFFFF"), HEADER_FILL, WRAP
        ws.column_dimensions[get_column_letter(col)].width = width
    for r, row in enumerate(rows, start=2):
        ws.append([v[0] if isinstance(v, tuple) else v for v in row])
        for col, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col)
            cell.alignment = WRAP
            if isinstance(v, tuple):  # (text, url)
                cell.hyperlink, cell.font = v[1], LINK
            if col > len(header) - review_cols:
                cell.fill = REVIEW_FILL
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions


def _evidence(item: dict) -> str:
    return "\n".join(f"{e['field']}: «{e['quote']}»" for e in item["evidence"])


def write(path: Path, store: Store | None = None) -> dict:
    batch = load(store or Store())
    fires = incidents(batch.reports)
    in_scope = [i for i in fires if i.scope != WAR]
    war = [i for i in fires if i.scope == WAR]
    articles = []
    for prefix, group in (("F", in_scope), ("W", war)):
        for n, i in enumerate(group, start=1):
            articles += [[f"{prefix}{n}", r.outlet, r.published.isoformat(), (r.headline, r.url), r.item["summary"],
                          _evidence(r.item)] for r in i.reports]
    set_aside = sorted(([(x["article"].get("published_at") or "")[:10], x["article"]["outlet"],
                         (x["article"]["headline"], x["article"]["url"]), SET_ASIDE_AS.get(x["scope"], x["scope"]),
                         x["item"]["summary"], round(x["item"]["confidence"], 2)] for x in batch.set_aside),
                       key=lambda row: row[0], reverse=True)
    failed = [[(a["article"].get("published_at") or "")[:10], a["article"]["outlet"],
               (a["article"]["headline"], a["article"]["url"]), a["status"]] for a in batch.failed]

    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    readme.sheet_view.showGridLines = False
    merged = len(batch.reports) - len(fires)
    lines = [
        ("Fire incidents in Lebanon — classification for review", True),
        ("", False),
        ((f"{len(batch.reports)} reports of fires in Lebanon became {len(fires)} fires once articles about the same fire "
          f"were merged ({merged} merged). {len(in_scope)} are on 'Fires in Lebanon'; {len(war)} war-related fires are "
          f"logged on their own sheet and will not be shown. Set aside: {len(set_aside)} (not a fire, abroad, background, "
          f"totals). Not classified: {len(failed)}."), False),
        ("", False),
        (("One row per fire, answering nine questions: date, location, the site that witnessed the fire, what burned, "
          "the hazard (what started or fed it), the origin (accident, deliberate, war), the damage, the gravity of the "
          "damage and the responders."), False),
        (("Under each answer: '— reported' means the article says it; '— deduced …' means the computer worked it out "
          "from something the article does say, and names what; 'not established' means neither. A deduction is made "
          "from reported facts only, and is never counted as reported."), False),
        (("Where a cell lists several values, the first is the main one and is the one counted. For the site, it is "
          "where the fire started (or the place that burned most); '(also: …)' lists the other places it reached, "
          "damaged or threatened, or was next to — as the school beside the burning land."), False),
        (("Gravity of the damage, from the damage reported (the brief's scale): 4 deaths or several buildings · 3 "
          "injuries or displaced people · 2 damage, no casualties · 1 no damage."), False),
        ("", False),
        (("How to review: write OK or wrong in the yellow 'YOUR VERDICT' column and what it should be in 'YOUR "
          "CORRECTION'. 'Articles' lists every article behind each fire (F3 = row 3 of 'Fires in Lebanon', W2 = row 2 "
          "of 'War-related') with the quotes Claude based its answers on. 'Set aside' shows what Claude did not count. "
          "'Lists' shows every category, its Arabic cue words and what it lets us deduce: correct an entry, or add one "
          "on a new row."), False),
        (("War-related fires (strikes, shelling, munitions, fires set by an army) are kept but not shown. Claude marks "
          "them; a second check flags in 'Checks' any other fire whose article uses war words."), False),
        (("Same fire: two articles are merged when their places match (same village, or a village inside the same "
          "district), their dates are at most two days apart, and their sites agree."), False),
    ]
    if batch.stale:
        lines.append((f"{batch.stale} classifications made with older instructions are not shown.", False))
    for text, bold in lines:
        readme.append([text])
        cell = readme.cell(row=readme.max_row, column=1)
        cell.font, cell.alignment = Font(bold=bold, size=14 if bold else 11), WRAP
    readme.column_dimensions["A"].width = 120
    _sheet(wb, "Fires in Lebanon", FIRE_COLUMNS, [fire_row(n, i) for n, i in enumerate(in_scope, start=1)], review_cols=2)
    _sheet(wb, "War-related (not shown)", FIRE_COLUMNS, [fire_row(n, i) for n, i in enumerate(war, start=1)], review_cols=2)
    _sheet(wb, "Articles", ARTICLE_COLUMNS, articles)
    _sheet(wb, "Set aside", SET_ASIDE_COLUMNS, [[n, *row] for n, row in enumerate(set_aside, start=1)])
    _sheet(wb, "Lists", LIST_COLUMNS, list_rows(), review_cols=1)
    if failed:
        _sheet(wb, "Not classified", [("#", 5), ("Published", 11), ("Outlet", 11), ("Headline", 60), ("Why", 30)],
               [[n, *row] for n, row in enumerate(failed, start=1)])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"reports": len(batch.reports), "fires": len(in_scope), "war_related": len(war), "merged": merged,
            "set_aside": len(set_aside), "failed": len(failed), "stale": batch.stale, "path": str(path)}
